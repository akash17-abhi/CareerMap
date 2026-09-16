from __future__ import annotations

import io
import re
from pathlib import Path

import fitz
from docx import Document


# ============================================================
# SUPPORTED FILE TYPES
# ============================================================

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
}

SUPPORTED_CONTENT_TYPES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


# ============================================================
# EXCEPTIONS
# ============================================================


class DocumentParserError(ValueError):
    """Raised when a document cannot be safely parsed."""


# ============================================================
# TEXT CLEANING
# ============================================================


def clean_extracted_text(text: str) -> str:
    """
    Clean common PDF/DOCX extraction noise while preserving
    meaningful resume/JD content.

    This function intentionally performs conservative cleaning.
    It must not rewrite, summarize, or interpret the document.
    """
    if not text:
        return ""

    # Normalize line endings.
    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    # Remove zero-width characters that can interfere with matching.
    text = re.sub(r"[\u200b\u200c\u200d\ufeff]", "", text)

    # Normalize non-breaking spaces.
    text = text.replace("\u00a0", " ")

    # Normalize horizontal whitespace while preserving line boundaries.
    text = re.sub(r"[ \t]+", " ", text)

    # Collapse excessive blank lines to at most one blank line.
    text = re.sub(r"\n[ \t]*\n[ \t]*\n+", "\n\n", text)

    # Strip whitespace around individual lines.
    lines = [line.strip() for line in text.split("\n")]

    # Remove consecutive blank lines defensively.
    cleaned_lines: list[str] = []
    previous_blank = False

    for line in lines:
        if not line:
            if previous_blank:
                continue
            cleaned_lines.append("")
            previous_blank = True
            continue

        cleaned_lines.append(line)
        previous_blank = False

    return "\n".join(cleaned_lines).strip()


# ============================================================
# VALIDATION
# ============================================================


def validate_document(
    file_bytes: bytes,
    filename: str,
    content_type: str | None = None,
) -> str:
    """
    Validate a document before attempting extraction.

    Returns the normalized file extension.
    """
    if not file_bytes:
        raise DocumentParserError("The uploaded document is empty.")

    if not filename or not filename.strip():
        raise DocumentParserError("The uploaded document has no filename.")

    safe_filename = Path(filename.strip()).name
    extension = Path(safe_filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentParserError(
            "Unsupported document format. Please upload a PDF or DOCX file."
        )

    if content_type:
        normalized_content_type = content_type.split(";")[0].strip().lower()

        # Browsers and clients may provide an empty or generic content type.
        # Use the extension as the primary compatibility check and reject
        # only types that are clearly incompatible with PDF/DOCX uploads.
        clearly_incompatible_types = {
            "text/plain",
            "image/png",
            "image/jpeg",
            "image/webp",
            "application/zip",
        }

        if normalized_content_type in clearly_incompatible_types:
            raise DocumentParserError(
                "The uploaded file type does not match a supported "
                "PDF or DOCX document."
            )

    return extension


# ============================================================
# PDF EXTRACTION
# ============================================================


def extract_pdf_text(file_bytes: bytes) -> str:
    """
    Extract text from a PDF entirely in memory.
    No temporary file is created.
    """
    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as document:
            if document.page_count == 0:
                raise DocumentParserError("The PDF does not contain any pages.")

            page_text: list[str] = []

            for page in document:
                try:
                    page_text.append(page.get_text("text"))
                except Exception as exc:
                    raise DocumentParserError(
                        "Text could not be extracted from one of the PDF pages."
                    ) from exc

            extracted_text = "\n".join(page_text)

    except DocumentParserError:
        raise
    except Exception as exc:
        raise DocumentParserError("The PDF could not be opened or parsed.") from exc

    cleaned_text = clean_extracted_text(extracted_text)

    if not cleaned_text:
        raise DocumentParserError(
            "No readable text was found in the PDF. "
            "The file may be scanned or image-only."
        )

    return cleaned_text


# ============================================================
# DOCX EXTRACTION
# ============================================================


def extract_docx_text(file_bytes: bytes) -> str:
    """
    Extract text from a DOCX document entirely in memory.

    Both normal paragraphs and table-cell text are extracted because
    resumes frequently store important sections such as Skills,
    Education, Experience, or Projects inside tables.
    """
    try:
        document = Document(io.BytesIO(file_bytes))
    except Exception as exc:
        raise DocumentParserError(
            "The DOCX file could not be opened or parsed."
        ) from exc

    text_parts: list[str] = []

    # --------------------------------------------------------
    # Normal paragraphs
    # --------------------------------------------------------
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            text_parts.append(text)

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------
    # Important for resume templates where Skills/Projects/etc.
    # may be stored inside a table instead of document.paragraphs.
    for table in document.tables:
        for row in table.rows:
            cell_texts: list[str] = []

            for cell in row.cells:
                paragraphs = [
                    paragraph.text.strip()
                    for paragraph in cell.paragraphs
                    if paragraph.text.strip()
                ]

                if paragraphs:
                    cell_text = " ".join(paragraphs).strip()
                    if cell_text:
                        cell_texts.append(cell_text)

            if cell_texts:
                text_parts.append(" | ".join(cell_texts))

    extracted_text = "\n".join(text_parts)
    cleaned_text = clean_extracted_text(extracted_text)

    if not cleaned_text:
        raise DocumentParserError(
            "No readable text was found in the DOCX file."
        )

    return cleaned_text


# ============================================================
# MAIN PARSER
# ============================================================


def extract_document_text(
    file_bytes: bytes,
    filename: str,
    content_type: str | None = None,
) -> str:
    """
    Extract clean text from a supported PDF or DOCX document.

    Processing is fully in memory and does not persist the uploaded
    document.

    Supported formats:
        - PDF
        - DOCX
    """
    extension = validate_document(
        file_bytes=file_bytes,
        filename=filename,
        content_type=content_type,
    )

    if extension == ".pdf":
        return extract_pdf_text(file_bytes)

    if extension == ".docx":
        return extract_docx_text(file_bytes)

    # This should never be reached because validate_document()
    # checks the supported extensions.
    raise DocumentParserError("Unsupported document format.")


# ============================================================
# MINIMUM CONTENT VALIDATION
# ============================================================


def validate_extracted_text(
    text: str,
    minimum_characters: int = 80,
) -> str:
    """
    Validate extracted text before sending it to downstream
    local analysis services.
    """
    cleaned_text = clean_extracted_text(text)

    if len(cleaned_text) < minimum_characters:
        raise DocumentParserError(
            "The document contains too little readable text for reliable analysis."
        )

    return cleaned_text
